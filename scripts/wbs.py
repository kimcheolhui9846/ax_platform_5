"""WBS(docs/WBS.xlsx) 고치기. 엑셀에서 직접 고친 칸은 건드리지 않고, 지정한 행·칸만 바꾼다.

실행(저장소 루트):
  uv run --no-project --with openpyxl python scripts/wbs.py list
  uv run --no-project --with openpyxl python scripts/wbs.py set 4.6 --status 완료 --end 2026-10-18 --pct 1 --output "PR #17"
  uv run --no-project --with openpyxl python scripts/wbs.py add 4.8 "작업 이름" --after 4.7 --owner 개발 --status 예정
  uv run --no-project --with openpyxl python scripts/wbs.py date 2026-10-18   (2행의 기준일)

단계 행(ID에 점이 없는 행)의 진척률은 하위 작업 평균 수식이고, '요약' 시트도 수식이라 따로 고치지 않는다.
"""
import argparse
import re
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Side

WBS = Path(__file__).resolve().parents[1] / "docs" / "WBS.xlsx"
SHEET = "WBS"
HEAD_ROW = 4
COLUMNS = {"id": 1, "phase": 2, "task": 3, "detail": 4, "owner": 5, "start": 6, "end": 7, "status": 8, "pct": 9, "output": 10, "memo": 11}
STATUSES = ("완료", "진행 중", "예정", "보류")
OWNERS = ("개발", "디자인", "공동")
_thin = Side(style="thin", color="D0D0D0")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def find_row(ws, task_id: str) -> int:
    for r in range(HEAD_ROW + 1, ws.max_row + 1):
        if str(ws.cell(r, COLUMNS["id"]).value) == task_id:
            return r
    sys.exit(f"없는 ID입니다: {task_id}")


def last_row(ws) -> int:
    return max(r for r in range(HEAD_ROW, ws.max_row + 1) if ws.cell(r, COLUMNS["id"]).value is not None)


def fix_formulas(ws) -> None:
    """행을 끼운 뒤 단계 행 평균 수식·자동 필터·검증·조건부 서식 범위를 다시 맞춘다."""
    end = last_row(ws)
    ids = {r: str(ws.cell(r, 1).value) for r in range(HEAD_ROW + 1, end + 1)}
    for r, rid in ids.items():
        if "." not in rid:
            children = [cr for cr, cid in ids.items() if cid.startswith(rid + ".")]
            if children:
                ws.cell(r, COLUMNS["pct"], f"=IFERROR(AVERAGE(I{min(children)}:I{max(children)}),0)")
    ws.auto_filter.ref = f"A{HEAD_ROW}:K{end}"
    for dv in ws.data_validations.dataValidation:
        col = str(dv.sqref).split(":")[0].rstrip("0123456789")
        dv.sqref = f"{col}{HEAD_ROW + 1}:{col}{end}"
    rules = ws.conditional_formatting
    renamed = []
    for rng in list(rules):
        col = str(rng.sqref).split(":")[0].rstrip("0123456789")
        renamed.append((f"{col}{HEAD_ROW + 1}:{col}{end}", rng.rules))
    rules._cf_rules.clear()
    for ref, rule_list in renamed:
        for rule in rule_list:
            rules.add(ref, rule)
    summary = ws.parent["요약"]
    for row in summary.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.value = re.sub(r"(\$[BHI]\$)\d+:(\$[BHI]\$)\d+", rf"\g<1>{HEAD_ROW + 1}:\g<2>{end}", cell.value)


def cmd_list(ws, _args) -> None:
    for r in range(HEAD_ROW + 1, last_row(ws) + 1):
        v = [ws.cell(r, c).value for c in (1, 3, 8, 9, 7, 10)]
        print(" | ".join("" if x is None else str(x) for x in v))


def cmd_set(ws, args) -> None:
    r = find_row(ws, args.id)
    if args.status and args.status not in STATUSES:
        sys.exit(f"상태는 {', '.join(STATUSES)} 중 하나입니다.")
    for name in ("task", "detail", "owner", "start", "end", "status", "output", "memo"):
        value = getattr(args, name)
        if value is not None:
            ws.cell(r, COLUMNS[name], value or None)
    if args.pct is not None:
        ws.cell(r, COLUMNS["pct"], args.pct)


def cmd_add(ws, args) -> None:
    after = find_row(ws, args.after)
    r = after + 1
    ws.insert_rows(r)
    phase = ws.cell(after, COLUMNS["phase"]).value
    values = {"id": args.id, "phase": phase, "task": args.task, "detail": args.detail, "owner": args.owner,
              "start": args.start, "end": args.end, "status": args.status, "pct": args.pct if args.pct is not None else 0,
              "output": args.output, "memo": args.memo}
    for name, c in COLUMNS.items():
        cell = ws.cell(r, c, values[name] or (0 if name == "pct" else None))
        cell.border = BORDER
        cell.alignment = Alignment(vertical="center", wrap_text=c in (3, 4, 10, 11))
    ws.cell(r, COLUMNS["pct"]).number_format = "0%"
    fix_formulas(ws)


def cmd_date(ws, args) -> None:
    text = str(ws["A2"].value or "")
    ws["A2"] = re.sub(r"기준일 \d{4}-\d{2}-\d{2}", f"기준일 {args.date}", text)


def main() -> None:
    p = argparse.ArgumentParser(description="docs/WBS.xlsx 고치기")
    sub = p.add_subparsers(dest="cmd", required=True)
    sub.add_parser("list")
    for name in ("set", "add"):
        sp = sub.add_parser(name)
        sp.add_argument("id")
        if name == "add":
            sp.add_argument("task")
            sp.add_argument("--after", required=True)
        else:
            sp.add_argument("--task")
        sp.add_argument("--detail")
        sp.add_argument("--owner", choices=OWNERS)
        sp.add_argument("--start")
        sp.add_argument("--end")
        sp.add_argument("--status", choices=STATUSES)
        sp.add_argument("--pct", type=float)
        sp.add_argument("--output")
        sp.add_argument("--memo")
    sd = sub.add_parser("date")
    sd.add_argument("date")
    args = p.parse_args()

    wb = load_workbook(WBS)
    ws = wb[SHEET]
    {"list": cmd_list, "set": cmd_set, "add": cmd_add, "date": cmd_date}[args.cmd](ws, args)
    if args.cmd != "list":
        wb.save(WBS)
        print(f"저장: {WBS.relative_to(WBS.parents[1])}")


if __name__ == "__main__":
    main()
